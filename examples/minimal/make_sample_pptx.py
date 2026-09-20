# -*- coding: utf-8 -*-
"""生成最小样例 PPTX（供 examples/minimal 端到端演练）

刻意覆盖流水线上的关键情形，这样"跑通"才有验证价值：
  · 一个跨多页、面积极小的角标图        -> 验证自动排除
  · 一个组合（group）内的图片            -> 验证递归与坐标换算
  · 宽图 / 方图 / 竖图三种比例共存      -> 验证定值画框（比例失配是常态）
  · 一个表格 + 参数行文本（含对齐空格） -> 验证探针计数与参数行切分（坑 4）
  · 无占位符标题的散文本框              -> 复现"抽取只得到无层级散文本"的真实处境

素材全部由脚本现场绘制（几何图形 + ASCII 标签），不含任何外部资源与真实项目信息。

用法：
    python make_sample_pptx.py [输出路径]        # 默认同目录 sample.pptx
"""
import os
import sys
import shutil
import tempfile

from pptx import Presentation
from pptx.util import Inches, Pt
from PIL import Image, ImageDraw

SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5


def make_image(path, size, rgb, label):
    """画一张便于肉眼判断比例与裁切的测试图：对角斜纹 + 边框 + ASCII 标签。"""
    w, h = size
    im = Image.new("RGB", (w, h), rgb)
    d = ImageDraw.Draw(im)
    step = max(24, int(min(w, h) / 12))
    for x in range(-h, w + h, step):
        d.line([(x, 0), (x + h, h)], fill=(255, 255, 255), width=max(2, step // 8))
    d.rectangle([0, 0, w - 1, h - 1], outline=(255, 255, 255), width=max(6, step // 4))
    d.text((24, 24), label, fill=(20, 20, 20))
    d.text((24, 24 + 18), "%dx%d" % (w, h), fill=(20, 20, 20))
    im.save(path, "PNG")


def textbox(slide, x, y, w, h, text, size, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    # 逐行建段：换行必须是独立段落，否则抽取阶段会把多行并成一段（真实 PPT 同理）
    for i, line in enumerate(text.split("\n")):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = para.add_run()
        run.text = line
        run.font.size = Pt(size)
    return tb


def picture(slide, path, x, y, w, h):
    return slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))


def build(out_path):
    assets = tempfile.mkdtemp(prefix="isc_sample_assets_")
    try:
        wide = os.path.join(assets, "wide.png")
        square = os.path.join(assets, "square.png")
        tall = os.path.join(assets, "tall.png")
        wide16 = os.path.join(assets, "wide16.png")
        tall9 = os.path.join(assets, "tall9.png")
        badge = os.path.join(assets, "badge.png")
        make_image(wide, (1800, 600), (46, 88, 148), "WIDE 3:1")
        make_image(square, (1200, 1200), (34, 128, 108), "SQUARE 1:1")
        make_image(tall, (700, 1400), (148, 78, 46), "TALL 1:2")
        make_image(wide16, (1600, 900), (96, 62, 148), "WIDE 16:9")
        make_image(tall9, (900, 1600), (140, 48, 88), "TALL 9:16")
        make_image(badge, (400, 200), (200, 200, 200), "FOOTER")

        prs = Presentation()
        prs.slide_width = Inches(SLIDE_W_IN)
        prs.slide_height = Inches(SLIDE_H_IN)
        blank = prs.slide_layouts[6]

        # --- P1 封面：大宽图（占版面 >15%，应进 1600 档） ---
        s = prs.slides.add_slide(blank)
        textbox(s, 0.9, 1.1, 11.5, 1.0, "Sample Deck", 40)
        textbox(s, 0.9, 2.2, 11.5, 0.6, "Minimal fixture for the extract / build / verify pipeline", 16)
        picture(s, wide, 0.8, 3.2, 11.7, 3.9)

        # --- P2 概述：方图 ---
        s = prs.slides.add_slide(blank)
        textbox(s, 0.9, 0.9, 8.0, 0.8, "Overview", 28)
        textbox(s, 0.9, 1.9, 6.0, 2.4,
                "Body copy taken verbatim from the source deck. "
                "This page exists to exercise paragraph extraction and per-page image ranking.", 14)
        picture(s, square, 7.4, 1.7, 5.0, 5.0)

        # --- P3 规格卡：竖图 + 参数行（含对齐空格，验证坑 4 的切分） ---
        s = prs.slides.add_slide(blank)
        textbox(s, 4.6, 0.9, 8.0, 0.8, "Spec Card", 28)
        textbox(s, 4.6, 1.9, 8.0, 0.5, "Spec: 120*45*0.6mm;   Density: 80kg/m3", 14)
        textbox(s, 4.6, 2.6, 8.0, 0.5, "Standard length: 3000mm;   Grade: A1", 14)
        textbox(s, 4.6, 3.3, 8.0, 2.4,
                "Bullet one: double-cavity structure for multi-stage insulation.\n"
                "Bullet two: pre-punched service holes, SI separation.\n"
                "Bullet three: shaped through holes for load transfer.", 14)
        picture(s, tall, 0.9, 1.4, 3.2, 6.4)

        # --- P4 对比：两种不同比例并排（验证定值画框） ---
        s = prs.slides.add_slide(blank)
        textbox(s, 0.9, 0.9, 11.0, 0.8, "Comparison", 28)
        picture(s, wide16, 0.9, 2.0, 5.5, 3.1)
        picture(s, tall9, 7.3, 2.0, 2.6, 4.6)

        # --- P5 组合：验证 group 递归与坐标换算 ---
        s = prs.slides.add_slide(blank)
        textbox(s, 0.9, 0.9, 11.0, 0.8, "Grouped Layout", 28)
        try:
            grp = s.shapes.add_group_shape()
            grp.shapes.add_picture(square, Inches(1.2), Inches(2.2), Inches(3.0), Inches(3.0))
            tb = grp.shapes.add_textbox(Inches(4.6), Inches(3.2), Inches(6.0), Inches(1.0))
            tb.text_frame.text = "Inside a group"
            grp.left, grp.top = Inches(0.9), Inches(1.9)
            grp.width, grp.height = Inches(10.5), Inches(4.0)
        except Exception:
            picture(s, square, 1.2, 2.2, 3.0, 3.0)   # 旧版 python-pptx 无组合 API 时降级
            textbox(s, 4.6, 3.2, 6.0, 1.0, "Inside a group (fallback)", 14)

        # --- P6 表格页：验证探针的表格计数 ---
        s = prs.slides.add_slide(blank)
        textbox(s, 0.9, 0.9, 11.0, 0.8, "Table Page", 28)
        shape = s.shapes.add_table(3, 3, Inches(0.9), Inches(2.2), Inches(7.5), Inches(2.0))
        for r in range(3):
            for c in range(3):
                shape.table.cell(r, c).text = "R%dC%d" % (r + 1, c + 1)
        textbox(s, 0.9, 4.8, 11.0, 1.0, "Tables must be reported by the probe (rebuilt as HTML, not a screenshot).", 14)

        # --- P7 / P8 续页（让角标出现页数达到阈值） ---
        for i in (7, 8):
            s = prs.slides.add_slide(blank)
            textbox(s, 0.9, 0.9, 11.0, 0.8, "Continuation %d" % i, 28)
            textbox(s, 0.9, 2.0, 8.0, 2.0,
                    "Filler page so the footer badge appears on enough pages "
                    "to trigger the recurring-decoration rule.", 14)

        # --- 角标：逐页右下角，面积极小（应被自动排除） ---
        for idx, s in enumerate(prs.slides, 1):
            picture(s, badge, SLIDE_W_IN - 1.1, SLIDE_H_IN - 0.75, 0.8, 0.4)

        prs.save(out_path)
        return len(prs.slides)
    finally:
        shutil.rmtree(assets, ignore_errors=True)


def main():
    here = os.path.dirname(os.path.abspath(__file__))
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(here, "sample.pptx")
    n = build(out)
    print("已生成: %s (%d 页, %.1f KB)" % (out, n, os.path.getsize(out) / 1024.0))
    return 0


if __name__ == "__main__":
    sys.exit(main())
